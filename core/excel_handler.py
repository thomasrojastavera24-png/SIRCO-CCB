"""
SIRCO v2.0 - Lectura y escritura del Excel de trabajo
Acepta rutas locales (Path) o archivos subidos (BytesIO).
"""
import io
import openpyxl
import shutil
from pathlib import Path
from datetime import datetime
from typing import Optional, Union

from core.config import (
    EXCEL_LOCAL_PATH, REPORTES_DIR,
    COL_NOMBRE, COL_NIT, COL_CAMARA, COL_FECHA,
    COL_TRAMITE, COL_OBS,
)

# Tipo que acepta Path o BytesIO
FuenteExcel = Union[Path, io.BytesIO]


def _detectar_hoja(wb: openpyxl.Workbook) -> str:
    """Detecta la hoja principal por nombre."""
    for name in wb.sheetnames:
        n = name.lower()
        if "revision" in n or "revisión" in n or "sociedad" in n:
            return name
    return wb.sheetnames[0]


def _abrir_workbook(fuente: FuenteExcel, data_only: bool = True) -> openpyxl.Workbook:
    """Abre un workbook desde Path o BytesIO."""
    if isinstance(fuente, io.BytesIO):
        fuente.seek(0)
        return openpyxl.load_workbook(fuente, data_only=data_only)
    return openpyxl.load_workbook(fuente, data_only=data_only)


def leer_sociedades_bogota(fuente: Optional[FuenteExcel] = None) -> list[dict]:
    """
    Lee el Excel y devuelve sociedades con Cámara = Bogotá.
    Cada elemento: {fila, nombre, nit_raw, nit_base, dv, camara, fecha, tramite, obs}
    """
    fuente = fuente or EXCEL_LOCAL_PATH
    wb = _abrir_workbook(fuente)
    ws = wb[_detectar_hoja(wb)]

    sociedades = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        nombre = str(row[COL_NOMBRE].value or "").strip()
        nit_raw = str(row[COL_NIT].value or "").strip()
        camara = str(row[COL_CAMARA].value or "").strip()
        fecha = str(row[COL_FECHA].value or "").strip()
        tramite = str(row[COL_TRAMITE].value or "").strip()
        obs = str(row[COL_OBS].value or "").strip()
        fila = row[COL_NOMBRE].row

        if nombre and "bogot" in camara.lower():
            nit_clean = nit_raw.replace(".", "").replace(" ", "").replace("-", "")
            nit_base = nit_clean[:-1] if len(nit_clean) >= 2 else nit_clean
            dv = nit_clean[-1] if len(nit_clean) >= 2 else ""
            sociedades.append({
                "fila": fila,
                "nombre": nombre,
                "nit_raw": nit_raw,
                "nit_base": nit_base,
                "dv": dv,
                "camara": camara,
                "fecha": fecha,
                "tramite": tramite,
                "obs": obs,
            })
    return sociedades


def leer_todas_las_sociedades(fuente: Optional[FuenteExcel] = None) -> list[dict]:
    """Lee todas las filas del Excel sin filtrar por cámara."""
    fuente = fuente or EXCEL_LOCAL_PATH
    wb = _abrir_workbook(fuente)
    ws = wb[_detectar_hoja(wb)]

    sociedades = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        nombre = str(row[COL_NOMBRE].value or "").strip()
        if not nombre:
            continue
        nit_raw = str(row[COL_NIT].value or "").strip()
        camara = str(row[COL_CAMARA].value or "").strip()
        fecha = str(row[COL_FECHA].value or "").strip()
        tramite = str(row[COL_TRAMITE].value or "").strip()
        obs = str(row[COL_OBS].value or "").strip()
        fila = row[COL_NOMBRE].row

        nit_clean = nit_raw.replace(".", "").replace(" ", "").replace("-", "")
        nit_base = nit_clean[:-1] if len(nit_clean) >= 2 else nit_clean
        dv = nit_clean[-1] if len(nit_clean) >= 2 else ""

        sociedades.append({
            "fila": fila,
            "nombre": nombre,
            "nit_raw": nit_raw,
            "nit_base": nit_base,
            "dv": dv,
            "camara": camara,
            "fecha": fecha,
            "tramite": tramite,
            "obs": obs,
        })
    return sociedades


def guardar_resultado_en_memoria(
    fuente: FuenteExcel,
    fila: int,
    tramite: str,
    obs: str,
    fecha: Optional[str] = None,
) -> io.BytesIO:
    """
    Actualiza una fila en el workbook y devuelve el archivo en memoria (BytesIO).
    No escribe en disco — compatible con Streamlit Cloud.
    """
    fecha = fecha or datetime.today().strftime("%d/%m/%Y")
    wb = _abrir_workbook(fuente, data_only=False)
    ws = wb[_detectar_hoja(wb)]

    ws.cell(row=fila, column=COL_FECHA + 1).value   = fecha
    ws.cell(row=fila, column=COL_TRAMITE + 1).value = tramite
    ws.cell(row=fila, column=COL_OBS + 1).value     = obs

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def guardar_resultados_lote_en_memoria(
    fuente: FuenteExcel,
    resultados: list[dict],
    fecha: Optional[str] = None,
) -> io.BytesIO:
    """
    Actualiza múltiples filas y devuelve el archivo en memoria.
    resultados: lista de {fila, tramite, obs}
    """
    fecha = fecha or datetime.today().strftime("%d/%m/%Y")
    wb = _abrir_workbook(fuente, data_only=False)
    ws = wb[_detectar_hoja(wb)]

    for r in resultados:
        fila = r["fila"]
        ws.cell(row=fila, column=COL_FECHA + 1).value   = fecha
        ws.cell(row=fila, column=COL_TRAMITE + 1).value = r["tramite"]
        ws.cell(row=fila, column=COL_OBS + 1).value     = r.get("obs", "")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
