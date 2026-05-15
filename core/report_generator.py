"""
SIRCO v2.0 - Generador de reportes
Genera reporte de trazabilidad en .txt y .docx.
Compatible con Streamlit Cloud: devuelve bytes en memoria, no escribe en disco.
"""
import io
from datetime import datetime
from typing import Optional

from core.config import PERIODO_DESDE, PERIODO_HASTA, CCB_URL


# ─────────────────────────────────────────────
# Reporte TXT (en memoria)
# ─────────────────────────────────────────────

def generar_reporte_txt_bytes(
    sociedades: list[dict],
    fecha_tag: Optional[str] = None,
) -> bytes:
    """
    Genera el reporte de trazabilidad como bytes UTF-8.
    No escribe en disco — compatible con st.download_button.
    """
    hoy = datetime.today().strftime("%d/%m/%Y")
    fecha_tag = fecha_tag or datetime.today().strftime("%d%m%Y")

    con_tramite = [s for s in sociedades if s.get("tramite", "").upper() == "SI"]
    sin_tramite = [s for s in sociedades if s.get("tramite", "").upper() != "SI"]
    incidencias = [
        s for s in sociedades
        if s.get("resultado_ccb") and getattr(s["resultado_ccb"], "incidencia", "")
    ]

    SEP = "=" * 80
    lineas = [
        SEP,
        "REPORTE DE TRAZABILIDAD - SEGUIMIENTO REGISTRAL CCB",
        SEP,
        f"Fecha de revisión:  {hoy}",
        f"Fuente consultada:  {CCB_URL}",
        f"Período evaluado:   {PERIODO_DESDE} - {PERIODO_HASTA}",
        "",
        "ALCANCE: Sociedades con Cámara de Comercio = Bogotá",
        "MÉTODO:  Búsqueda avanzada por NIT en plataforma CCB en línea",
        SEP,
        "",
        "RESUMEN EJECUTIVO",
        "-" * 40,
        f"Total sociedades revisadas:         {len(sociedades)}",
        f"Con trámites visibles (SI):          {len(con_tramite)}",
        f"Sin trámites visibles (NO):          {len(sin_tramite)}",
        "",
        SEP,
        "SOCIEDADES CON TRÁMITES EN 2026 (SI)",
        SEP,
    ]

    if con_tramite:
        for i, s in enumerate(con_tramite, 1):
            obs = s.get("obs", "")
            lineas += [f"\n{i}. {s['nombre']}", f"   NIT: {s['nit_raw']}", f"   Observaciones: {obs}"]
    else:
        lineas.append("  (Ninguna)")

    lineas += ["", SEP, "SOCIEDADES SIN TRÁMITES EN 2026 (NO)", SEP]
    for i, s in enumerate(sin_tramite, 1):
        lineas.append(f"  {i:2d}. {s['nombre']} | NIT: {s['nit_raw']}")

    lineas += ["", SEP, "INCIDENCIAS / NOVEDADES", SEP]
    if incidencias:
        for s in incidencias:
            lineas.append(f"  - {s['resultado_ccb'].incidencia}")
    else:
        lineas.append("  (Sin incidencias)")

    lineas += [
        "", SEP,
        f"Reporte generado por SIRCO v2.0 — {hoy}",
        SEP,
    ]

    return "\n".join(lineas).encode("utf-8")


# ─────────────────────────────────────────────
# Reporte Word (.docx) en memoria
# ─────────────────────────────────────────────

def generar_reporte_docx_bytes(
    sociedades: list[dict],
    fecha_tag: Optional[str] = None,
) -> bytes:
    """
    Genera el reporte de trazabilidad como bytes .docx.
    No escribe en disco — compatible con st.download_button.
    """
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    hoy = datetime.today().strftime("%d/%m/%Y")
    fecha_tag = fecha_tag or datetime.today().strftime("%d%m%Y")

    con_tramite = [s for s in sociedades if s.get("tramite", "").upper() == "SI"]
    sin_tramite = [s for s in sociedades if s.get("tramite", "").upper() != "SI"]
    incidencias = [
        s for s in sociedades
        if s.get("resultado_ccb") and getattr(s["resultado_ccb"], "incidencia", "")
    ]

    doc = Document()

    titulo = doc.add_heading("REPORTE DE TRAZABILIDAD — SEGUIMIENTO REGISTRAL CCB", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Fecha de revisión: {hoy}")
    doc.add_paragraph(f"Período evaluado: {PERIODO_DESDE} – {PERIODO_HASTA}")
    doc.add_paragraph(f"Fuente: {CCB_URL}")
    doc.add_paragraph("Alcance: Sociedades con Cámara de Comercio = Bogotá")
    doc.add_paragraph("")

    doc.add_heading("Resumen ejecutivo", level=2)
    doc.add_paragraph(f"Total revisadas: {len(sociedades)}")
    doc.add_paragraph(f"Con trámites (SI): {len(con_tramite)}")
    doc.add_paragraph(f"Sin trámites (NO): {len(sin_tramite)}")

    doc.add_heading("Sociedades con trámites en 2026 (SI)", level=2)
    if con_tramite:
        for s in con_tramite:
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(f"{s['nombre']}  |  NIT: {s['nit_raw']}")
            run.bold = True
            obs = s.get("obs", "")
            if obs:
                doc.add_paragraph(f"     {obs}")
    else:
        doc.add_paragraph("(Ninguna)")

    doc.add_heading("Sociedades sin trámites en 2026 (NO)", level=2)
    for s in sin_tramite:
        doc.add_paragraph(f"{s['nombre']}  |  NIT: {s['nit_raw']}", style="List Bullet")

    doc.add_heading("Incidencias / novedades", level=2)
    if incidencias:
        for s in incidencias:
            doc.add_paragraph(s["resultado_ccb"].incidencia, style="List Bullet")
    else:
        doc.add_paragraph("Sin incidencias.")

    doc.add_paragraph(f"\nReporte generado por SIRCO v2.0 — {hoy}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
